from pathlib import Path

from openpyxl import load_workbook

from app.services import supabase_client

DATA_PATH = Path(__file__).resolve().parent.parent.parent / "data" / "inquiries.xlsx"

# 엑셀 컬럼 순서와 내부에서 사용할 필드명 매핑
COLUMNS = [
    "inquiry_id",
    "employee_id",
    "inquiry_date",
    "inquiry_text",
    "language",
    "category",
    "flag",
    "urgency",
    "status",
    "final_result",
]

HEADER_ROW = 3  # 1행: 안내 타이틀, 2행: 공백, 3행: 헤더

_inquiries_cache = None


def _load_from_excel():
    wb = load_workbook(DATA_PATH, data_only=True)
    ws = wb["문의내역"]

    inquiries = []
    for row in ws.iter_rows(min_row=HEADER_ROW + 1, values_only=True):
        if row[0] is None:
            continue
        record = dict(zip(COLUMNS, row))
        if hasattr(record["inquiry_date"], "strftime"):
            record["inquiry_date"] = record["inquiry_date"].strftime("%Y-%m-%d")
        inquiries.append(record)
    return inquiries


def _load_from_supabase(client):
    result = client.table("inquiries").select("*").order("inquiry_id").execute()
    return result.data


def get_all_inquiries():
    """전체 문의 목록을 반환한다. Supabase가 설정되어 있으면 그쪽에서, 아니면
    로컬 엑셀 파일(data/inquiries.xlsx)에서 읽는다."""
    global _inquiries_cache
    if _inquiries_cache is None:
        client = supabase_client.get_client()
        _inquiries_cache = _load_from_supabase(client) if client else _load_from_excel()
    return _inquiries_cache


def get_inquiry_by_id(inquiry_id):
    for inquiry in get_all_inquiries():
        if inquiry["inquiry_id"] == inquiry_id:
            return inquiry
    return None


def get_inquiries_by_employee(employee_id):
    return [i for i in get_all_inquiries() if i["employee_id"] == employee_id]
