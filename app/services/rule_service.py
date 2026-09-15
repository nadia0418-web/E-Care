from pathlib import Path

from openpyxl import load_workbook

from app.services import supabase_client

# 가상 GHD 업무 판단 기준 데이터 (실제 ETNERS/고객사의 공식 규정 아님, 프로젝트 시연용)
DATA_PATH = Path(__file__).resolve().parent.parent.parent / "data" / "ghd_rules.xlsx"

# 엑셀 컬럼 순서와 내부에서 사용할 필드명 매핑
COLUMNS = [
    "rule_id",
    "category",
    "rule_name",
    "situation",
    "keywords",
    "handler",
    "hr_flag",
    "urgency",
    "reasoning",
    "note",
]

HEADER_ROW = 3  # 1행: 안내 타이틀, 2행: 공백, 3행: 헤더

_rules_cache = None


def _load_from_excel():
    wb = load_workbook(DATA_PATH, data_only=True)
    ws = wb["가상GHD판단기준"]

    rows = []
    for row in ws.iter_rows(min_row=HEADER_ROW + 1, values_only=True):
        if row[0] is None:
            continue
        rows.append(dict(zip(COLUMNS, row)))
    return rows


def _load_from_supabase(client):
    result = client.table("rules").select("*").order("rule_id").execute()
    return result.data


def _normalize(record):
    record["hr_flag"] = bool(record["hr_flag"])
    record["keywords"] = [
        kw.strip() for kw in (record["keywords"] or "").split(",") if kw.strip()
    ]
    return record


def get_all_rules():
    """전체 가상 GHD 업무 판단 기준을 반환한다. Supabase가 설정되어 있으면 그쪽에서,
    아니면 로컬 엑셀 파일(data/ghd_rules.xlsx)에서 읽는다."""
    global _rules_cache
    if _rules_cache is None:
        client = supabase_client.get_client()
        rows = _load_from_supabase(client) if client else _load_from_excel()
        _rules_cache = [_normalize(dict(r)) for r in rows]
    return _rules_cache


def get_rule_by_id(rule_id):
    for rule in get_all_rules():
        if rule["rule_id"] == rule_id:
            return rule
    return None


def get_rules_by_category(category):
    return [r for r in get_all_rules() if r["category"] == category]
