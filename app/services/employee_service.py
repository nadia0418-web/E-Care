import json
import re
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from app.services import supabase_client

DATA_PATH = Path(__file__).resolve().parent.parent.parent / "data" / "employees.xlsx"

# 엑셀 컬럼 순서와 내부에서 사용할 필드명 매핑
COLUMNS = [
    "employee_id",
    "name",
    "nationality",
    "client_company",
    "position",
    "start_date",
    "employment_period",
    "family_accompanied",
    "visa_type",
    "special_notes",
]

# 신규 등록(수기 입력/엑셀 업로드) 시 실제로 입력받는 필드와 순서.
# employment_period는 더 이상 사용자가 입력하지 않는다 — 근무기간(재직기간)은
# 입사일을 기준으로 화면에서 매번 자동 계산해 보여준다 (compute_tenure_label 참고).
# COLUMNS에는 기존 엑셀 파일과의 컬럼 위치 호환을 위해 그대로 남겨두되, 값은
# 항상 빈 문자열로 저장한다.
UPLOAD_COLUMNS = ["name", "nationality", "client_company", "position", "start_date", "family_accompanied", "visa_type", "special_notes"]

# 정착 체크리스트 상태(JSON 문자열)는 신규 등록 폼/일괄 업로드 대상이 아니라
# 별도로 관리한다 (COLUMNS 바로 다음 컬럼에 저장).
SETTLEMENT_COLUMN = "settlement_checklist"
_ALL_COLUMNS = COLUMNS + [SETTLEMENT_COLUMN]

HEADER_ROW = 3  # 1행: 안내 타이틀, 2행: 공백, 3행: 헤더

_employees_cache = None


def _load_from_excel():
    wb = load_workbook(DATA_PATH, data_only=True)
    ws = wb["직원명단"]

    employees = []
    for row in ws.iter_rows(min_row=HEADER_ROW + 1, values_only=True):
        if row[0] is None:
            continue
        record = dict(zip(COLUMNS, row))
        record["family_accompanied"] = record["family_accompanied"] == "예"
        record["special_notes"] = record.get("special_notes") or ""
        if hasattr(record["start_date"], "strftime"):
            record["start_date"] = record["start_date"].strftime("%Y-%m-%d")
        settlement_raw = row[len(COLUMNS)] if len(row) > len(COLUMNS) else None
        record[SETTLEMENT_COLUMN] = settlement_raw or "{}"
        employees.append(record)
    return employees


def _load_from_supabase(client):
    result = client.table("employees").select("*").order("employee_id").execute()
    return result.data


def get_all_employees():
    """전체 직원 목록을 반환한다. Supabase가 설정되어 있으면 그쪽에서, 아니면
    로컬 엑셀 파일(data/employees.xlsx)에서 읽는다 (배포 환경에서는 Supabase,
    로컬 개발 환경에서는 엑셀을 그대로 사용할 수 있도록 둘 다 지원)."""
    global _employees_cache
    if _employees_cache is None:
        client = supabase_client.get_client()
        _employees_cache = _load_from_supabase(client) if client else _load_from_excel()
    return _employees_cache


def compute_tenure_label(start_date, today=None):
    """입사일(start_date)부터 오늘까지 경과한 재직기간을 "N년 M개월" 형태로 계산한다.
    완료된 개월 수만 센다(예: 입사 11개월+20일 차라면 "11개월", 반올림하지 않음).
    입사일이 없거나 미래 날짜면 "-"를 반환한다."""
    today = today or date.today()
    if isinstance(start_date, str):
        try:
            start_date = date.fromisoformat(start_date)
        except (TypeError, ValueError):
            return "-"
    if not start_date or start_date > today:
        return "-"

    months = (today.year - start_date.year) * 12 + (today.month - start_date.month)
    if today.day < start_date.day:
        months -= 1
    months = max(months, 0)

    years, rem_months = divmod(months, 12)
    if years and rem_months:
        return f"{years}년 {rem_months}개월"
    if years:
        return f"{years}년"
    return f"{rem_months}개월"


def get_filter_options(employees):
    """직원 목록 화면의 검색 필터(국적/소속회사/비자유형) 선택지를 중복 없이
    정렬해 반환한다."""
    def _distinct_sorted(values):
        return sorted({v for v in values if v})

    return {
        "nationalities": _distinct_sorted(e.get("nationality") for e in employees),
        "client_companies": _distinct_sorted(e.get("client_company") for e in employees),
        "visa_types": _distinct_sorted(e.get("visa_type") for e in employees),
    }


def get_employee_by_id(employee_id):
    for employee in get_all_employees():
        if employee["employee_id"] == employee_id:
            return employee
    return None


def _max_existing_num(existing):
    """기존 사번(EMP001 등)에서 가장 큰 숫자를 뽑아낸다. 다음 신규 사번 채번에 사용."""
    nums = []
    for e in existing:
        m = re.search(r"(\d+)$", str(e.get("employee_id") or ""))
        if m:
            nums.append(int(m.group(1)))
    return max(nums) if nums else 0


def _normalize_new_record(record):
    """폼 입력/엑셀 업로드에서 받은 값을 내부 표준 형태로 정리한다."""
    return {
        "name": (record.get("name") or "").strip(),
        "nationality": (record.get("nationality") or "").strip(),
        "client_company": (record.get("client_company") or "").strip(),
        "position": (record.get("position") or "").strip(),
        "start_date": (record.get("start_date") or "").strip() if isinstance(record.get("start_date"), str) else record.get("start_date"),
        "employment_period": (record.get("employment_period") or "").strip(),
        "family_accompanied": bool(record.get("family_accompanied")),
        "visa_type": (record.get("visa_type") or "").strip(),
        "special_notes": (record.get("special_notes") or "").strip(),
    }


def _append_to_excel(records):
    wb = load_workbook(DATA_PATH)
    ws = wb["직원명단"]
    next_row = ws.max_row + 1
    for offset, record in enumerate(records):
        row = next_row + offset
        start_date = record["start_date"]
        if isinstance(start_date, str) and start_date:
            try:
                start_date = date.fromisoformat(start_date)
            except ValueError:
                pass
        values = [
            record["employee_id"],
            record["name"],
            record["nationality"],
            record["client_company"],
            record["position"],
            start_date,
            record["employment_period"],
            "예" if record["family_accompanied"] else "아니오",
            record["visa_type"],
            record.get("special_notes", ""),
            record.get(SETTLEMENT_COLUMN, "{}"),
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(row=row, column=col, value=value)
    wb.save(DATA_PATH)


def _insert_to_supabase(client, records):
    payload = []
    for record in records:
        start_date = record["start_date"]
        if hasattr(start_date, "isoformat"):
            start_date = start_date.isoformat()
        payload.append({**record, "start_date": start_date})
    client.table("employees").upsert(payload).execute()


def _update_in_excel(employee_id, patch):
    wb = load_workbook(DATA_PATH)
    ws = wb["직원명단"]
    col_index = {name: i + 1 for i, name in enumerate(_ALL_COLUMNS)}
    for row in ws.iter_rows(min_row=HEADER_ROW + 1):
        if row[0].value == employee_id:
            for field, value in patch.items():
                col = col_index.get(field)
                if not col:
                    continue
                if field == "family_accompanied":
                    value = "예" if value else "아니오"
                elif field == "start_date" and isinstance(value, str) and value:
                    try:
                        value = date.fromisoformat(value)
                    except ValueError:
                        pass
                row[col - 1].value = value
            break
    wb.save(DATA_PATH)


def _update_in_supabase(client, employee_id, patch):
    client.table("employees").update(patch).eq("employee_id", employee_id).execute()


def update_employee(employee_id, patch):
    """기존 직원 레코드의 일부 필드를 수정한다 (직원 정보 수정, 정착 체크리스트 저장 등
    공용으로 사용). Supabase가 설정되어 있으면 그쪽에, 아니면 로컬 엑셀에 반영하고
    캐시를 비운다."""
    global _employees_cache

    client = supabase_client.get_client()
    if client:
        _update_in_supabase(client, employee_id, patch)
    else:
        _update_in_excel(employee_id, patch)

    _employees_cache = None


def add_employees(new_records):
    """신규 직원 레코드 목록을 추가한다 (수기 입력/엑셀 업로드 공용).
    각 레코드는 employee_id 없이 들어오며 여기서 사번을 자동 채번한다.
    Supabase가 설정되어 있으면 그쪽에, 아니면 로컬 엑셀에 반영하고, 캐시를
    비워 다음 조회부터 다른 화면(대시보드/Proactive Care 등)에도 바로 반영되게 한다."""
    global _employees_cache

    existing = get_all_employees()
    normalized = [_normalize_new_record(r) for r in new_records]
    normalized = [r for r in normalized if r["name"]]  # 이름이 없는 행은 건너뜀
    if not normalized:
        return []

    start_num = _max_existing_num(existing) + 1
    width = max(3, len(str(start_num + len(normalized) - 1)))
    for i, record in enumerate(normalized):
        record["employee_id"] = f"EMP{start_num + i:0{width}d}"
        record[SETTLEMENT_COLUMN] = "{}"  # 신규 입사자는 정착 체크리스트 미체크 상태로 시작

    client = supabase_client.get_client()
    if client:
        _insert_to_supabase(client, normalized)
    else:
        _append_to_excel(normalized)

    _employees_cache = None
    return normalized


def parse_upload_workbook(file_stream):
    """직원 데이터 일괄 업로드용 엑셀을 읽는다.

    기대 형식: 1행은 헤더(자동 무시), 2행부터 데이터. 컬럼 순서는
    이름, 국적, 소속회사, 직급, 입사일(YYYY-MM-DD), 가족동반(예/아니오), 비자유형, 특이사항(선택).
    근무기간(재직기간)은 입사일을 기준으로 자동 계산되므로 업로드 항목에 포함하지 않는다.
    이름이 비어 있는 행은 건너뛴다.
    """
    wb = load_workbook(file_stream, data_only=True)
    ws = wb.worksheets[0]

    records = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        raw = dict(zip(UPLOAD_COLUMNS, row))
        start_date = raw.get("start_date")
        if hasattr(start_date, "isoformat"):
            start_date = start_date.isoformat()
        raw["start_date"] = start_date
        raw["family_accompanied"] = str(raw.get("family_accompanied") or "").strip() == "예"
        records.append(raw)
    return records
